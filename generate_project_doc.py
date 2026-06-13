"""
Generates the formal Project Document for E-Bringgs Technologies as a .docx file.
Run:  python generate_project_doc.py
Output: E-Bringgs_Project_Document.docx
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import date

OUT = "E-Bringgs_Project_Document.docx"

# ---------- Styling helpers ----------
BRAND = RGBColor(0x0B, 0x5F, 0xFF)   # blue
DARK  = RGBColor(0x11, 0x18, 0x27)
GREY  = RGBColor(0x4B, 0x55, 0x63)


def set_cell_bg(cell, hex_color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_heading(doc, text, level=1, color=BRAND):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = color
        run.font.name = "Calibri"
    return h


def add_para(doc, text, bold=False, italic=False, size=11, color=DARK, align=None,
             space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = "Calibri"
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    return p


def add_bullets(doc, items, style="List Bullet"):
    for item in items:
        p = doc.add_paragraph(style=style)
        run = p.add_run(item)
        run.font.size = Pt(11)
        run.font.name = "Calibri"
        run.font.color.rgb = DARK
        p.paragraph_format.space_after = Pt(3)


def add_page_break(doc):
    doc.add_page_break()


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER

    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(11)
        run.font.name = "Calibri"
        set_cell_bg(hdr[i], "0B5FFF")
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(10.5)
            run.font.name = "Calibri"
            run.font.color.rgb = DARK
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = w
    return table


# ---------- Build document ----------
doc = Document()

# Default font
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

# Page margins
for section in doc.sections:
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

# ------------- COVER -------------
for _ in range(4):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("E-BRINGGS TECHNOLOGIES")
r.font.size = Pt(34)
r.font.bold = True
r.font.color.rgb = BRAND
r.font.name = "Calibri"

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Project Document")
r.font.size = Pt(22)
r.font.color.rgb = DARK
r.font.name = "Calibri"

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("A Unified EdTech, Services & Talent Marketplace Platform")
r.italic = True
r.font.size = Pt(13)
r.font.color.rgb = GREY
r.font.name = "Calibri"

for _ in range(8):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(f"Prepared: {date.today().strftime('%B %Y')}")
r.font.size = Pt(12)
r.font.color.rgb = GREY
r.font.name = "Calibri"

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Version 1.0  |  Confidential")
r.font.size = Pt(11)
r.font.color.rgb = GREY
r.font.name = "Calibri"

add_page_break(doc)

# ------------- TABLE OF CONTENTS -------------
add_heading(doc, "Table of Contents", level=1)
toc_items = [
    "1.  Executive Summary",
    "2.  Project Scope",
    "3.  Problem Statement",
    "4.  Project Objectives",
    "5.  Product Overview — How E-Bringgs Works",
    "6.  Services Offered",
    "7.  Target Audience & User Roles",
    "8.  System Architecture & Technology Stack",
    "9.  Key Features by Module",
    "10. Market Landscape & Competitive Positioning",
    "11. How E-Bringgs Changes the Trend",
    "12. Business & Revenue Model",
    "13. Implementation Status",
    "14. Future Roadmap",
    "15. Risks & Mitigation",
    "16. Conclusion",
]
for item in toc_items:
    p = doc.add_paragraph()
    r = p.add_run(item)
    r.font.size = Pt(11)
    r.font.color.rgb = DARK
    r.font.name = "Calibri"
    p.paragraph_format.space_after = Pt(3)

add_page_break(doc)

# ------------- 1. EXECUTIVE SUMMARY -------------
add_heading(doc, "1. Executive Summary", level=1)
add_para(doc,
    "E-Bringgs Technologies is a full-stack digital platform that unifies three pillars "
    "of the modern knowledge economy — structured technology learning, professional "
    "services, and freelance/consulting engagement — into a single, role-aware ecosystem. "
    "It is built for the African market (with Nigerian Naira pricing through Paystack) "
    "while being globally deployable.")
add_para(doc,
    "The platform serves three distinct user types — Students enrolled in cohort programs, "
    "Clients commissioning services and projects, and Administrators running the business — "
    "each with their own dashboard, workflows, and permissions. Live classes, mentorship, "
    "project tracking, payments, blogging, and AI-assisted tutoring all sit within one "
    "coherent experience.")
add_para(doc,
    "This document defines the scope of the product, the problems it solves, the services "
    "it delivers, the technology it is built on, and the way it intends to reshape how "
    "African talent is trained, deployed, and connected to opportunity.")

add_page_break(doc)

# ------------- 2. PROJECT SCOPE -------------
add_heading(doc, "2. Project Scope", level=1)

add_heading(doc, "2.1 In Scope", level=2)
add_bullets(doc, [
    "A public-facing marketing site — landing page, pricing, services catalog, blog, "
    "About / Contact / Terms / Privacy.",
    "User authentication with email verification, password reset, JWT access tokens "
    "and refresh tokens, and role-based access control (Admin, Student, Client).",
    "Student dashboard — enrolled courses, learning progress, downloadable resources, "
    "live class schedule, assignments, notifications.",
    "Client dashboard — project request and tracking, payment history, dedicated "
    "client area with sidebar navigation.",
    "Admin dashboard — user management, course management, blog management, service "
    "request review, payments overview, analytics, role assignment, live session "
    "management, cohort management.",
    "Course Management System — course creation, modules and lessons, cohort "
    "scheduling, assignment upload and review.",
    "Live Video Classroom — WebRTC-based video calls, screen sharing, in-class chat, "
    "instructor controls (mute-all, kick), attendance tracking.",
    "Blog System — markdown/rich-text editor, categories and tags, search, SEO "
    "metadata.",
    "Payments — Paystack integration (NGN), subscription and one-time payments, "
    "transaction storage, success/failure flows, webhook reconciliation.",
    "AI-assisted tutoring (Google Gemini) and WhatsApp notifications (Meta Cloud API) "
    "as growth and retention features.",
    "Production-grade security — Helmet, CORS, rate limiting, input validation, "
    "secure environment variable handling.",
])

add_heading(doc, "2.2 Out of Scope (Phase 1)", level=2)
add_bullets(doc, [
    "Native mobile applications (iOS/Android) — the platform is delivered as a "
    "responsive web app initially.",
    "On-platform peer-to-peer messaging between students.",
    "Marketplace for third-party instructors (the platform delivers its own curriculum "
    "in phase 1).",
    "Advanced learning analytics (deep skill graphs, predictive performance models) — "
    "deferred to a later phase.",
])

add_heading(doc, "2.3 Deliverables", level=2)
add_bullets(doc, [
    "Production-ready frontend (React 18 + TypeScript + Vite + Tailwind CSS v4).",
    "Production-ready backend (Node.js + Express + TypeScript + MongoDB).",
    "WebSocket signaling server for live video classrooms.",
    "Swagger / OpenAPI API documentation accessible at /api/docs.",
    "Deployment-ready environment configuration (development, staging, production).",
])

add_page_break(doc)

# ------------- 3. PROBLEM STATEMENT -------------
add_heading(doc, "3. Problem Statement", level=1)
add_para(doc,
    "Across Africa — and Nigeria in particular — three painful gaps continue to block "
    "the flow of talent into the global digital economy:")

add_heading(doc, "3.1 The Training Gap", level=2)
add_para(doc,
    "Most aspiring software developers, data analysts, and product designers do not have "
    "access to structured, mentor-led, cohort-based training in their own currency and at "
    "a price they can afford. Foreign bootcamps charge in dollars, local options often "
    "lack rigor, and free YouTube content is unstructured and isolating.")

add_heading(doc, "3.2 The Trust Gap (Services)", level=2)
add_para(doc,
    "Businesses that need software development, data analysis, research support, or UX "
    "work struggle to find a single trusted partner. They juggle freelancers from "
    "scattered platforms, with no shared dashboard, no transparent pricing, and no "
    "accountability for delivery. The result is wasted budget and missed deadlines.")

add_heading(doc, "3.3 The Continuity Gap (Talent → Opportunity)", level=2)
add_para(doc,
    "Even when students complete training, there is no smooth bridge from being a learner "
    "to being a paid contributor on real client work. Schools end at the certificate; "
    "marketplaces start at zero reputation. The talent pipeline leaks at exactly the point "
    "where it should be flowing fastest.")

add_heading(doc, "3.4 Operational Fragmentation", level=2)
add_para(doc,
    "Existing solutions stitch together five or more disconnected tools — a learning "
    "management system, a Zoom link, a WhatsApp group, a Stripe invoice, a Notion board. "
    "Students lose context, clients lose visibility, admins lose hours. There is no single "
    "source of truth for the relationship between a person, a programme, a payment, and a "
    "project.")

add_page_break(doc)

# ------------- 4. OBJECTIVES -------------
add_heading(doc, "4. Project Objectives", level=1)
add_bullets(doc, [
    "Deliver a single, role-aware platform where Students, Clients, and Administrators "
    "each see exactly the workflow they need — and nothing they don't.",
    "Make high-quality, mentor-led tech training accessible in Nigerian Naira through "
    "Paystack, with cohort programmes, mentorship, and consulting tiers.",
    "Provide a transparent services dashboard where clients can request, track, and pay "
    "for software development, data analysis, research support, and UX/Product work.",
    "Build a native live-class experience using WebRTC — no third-party Zoom dependency, "
    "with attendance tracking and instructor controls baked in.",
    "Bridge the talent-to-opportunity gap by giving top-performing students a path into "
    "client projects directly inside the same platform they trained on.",
    "Achieve production-readiness with secure authentication, rate limiting, input "
    "validation, structured logging, and full API documentation.",
])

add_page_break(doc)

# ------------- 5. PRODUCT OVERVIEW -------------
add_heading(doc, "5. Product Overview — How E-Bringgs Works", level=1)
add_para(doc,
    "E-Bringgs is delivered as a responsive web application accessible from any modern "
    "browser. A visitor lands on the public marketing site and is funnelled — based on "
    "intent — into one of three experiences:")

add_heading(doc, "5.1 The Student Journey", level=2)
add_para(doc,
    "A prospective learner browses the Pricing and Courses pages, registers as a Student, "
    "verifies their email, and pays in NGN through Paystack. They are redirected to the "
    "Student Dashboard which shows enrolled courses, modules, downloadable resources, "
    "their cohort schedule, assignments, and any live session for the day. When a class "
    "is live, the Classroom page hosts a WebRTC video room with chat, screen-sharing, "
    "mute/unmute and instructor controls.")

add_heading(doc, "5.2 The Client Journey", level=2)
add_para(doc,
    "A business visitor browses the Services dashboard — Software Development, Data "
    "Analysis, Research Support, UX/Product — submits a service inquiry, registers as a "
    "Client, and is taken into the Client area. From the Client sidebar they can view "
    "active projects, drill into a project's status, see deliverables, and review payment "
    "history. All under a single dashboard.")

add_heading(doc, "5.3 The Admin Journey", level=2)
add_para(doc,
    "Administrators log in to a dedicated dark-themed admin layout. They manage users "
    "and roles, create and update courses (including individual modules), publish blog "
    "posts, schedule live sessions, manage cohorts, review and assign service requests, "
    "review submitted assignments, monitor payments, and consult dashboard analytics.")

add_heading(doc, "5.4 The Underlying Mechanics", level=2)
add_para(doc,
    "Behind the scenes, the React frontend communicates with the Express backend through "
    "a single Axios instance with auto-refreshing JWTs. State is persisted client-side "
    "with Zustand. Live video sessions open a WebSocket connection to a signaling server "
    "on the /ws path which orchestrates WebRTC offers, answers, ICE candidates, chat, "
    "attendance, and instructor commands. Payments are initialised through Paystack and "
    "verified via webhook back into the database. Files are uploaded through a Multer-"
    "backed endpoint with MIME allowlists and size limits.")

add_page_break(doc)

# ------------- 6. SERVICES OFFERED -------------
add_heading(doc, "6. Services Offered", level=1)
add_para(doc,
    "E-Bringgs delivers value through three product lines, each with its own pricing tier "
    "and dashboard experience.")

add_heading(doc, "6.1 Learning Programmes", level=2)
add_bullets(doc, [
    "Cohort programmes — structured, time-boxed group learning with live classes.",
    "Mentorship tier — 1:1 sessions with industry mentors.",
    "Self-paced course library — modules and lessons accessible on demand.",
    "Live video classroom — WebRTC-based, with chat, screen share, attendance.",
    "Assignment submission and instructor review.",
])

add_heading(doc, "6.2 Professional Services", level=2)
add_bullets(doc, [
    "Software Development — web, mobile, and full-stack engineering for clients.",
    "Data Analysis — dashboards, reports, and analytical work.",
    "Research Support — desk research, literature reviews, structured analysis.",
    "UX / Product Design — research, wireframing, design systems, usability.",
])

add_heading(doc, "6.3 Consulting & Freelance Engagements", level=2)
add_bullets(doc, [
    "Strategic consulting on product, engineering, and data initiatives.",
    "Short-term freelance engagements channelled through the Client dashboard.",
    "Retainer-style relationships with structured deliverables and visibility.",
])

add_heading(doc, "6.4 Pricing Snapshot", level=2)
add_table(
    doc,
    headers=["Tier", "Audience", "Billing Model", "Currency"],
    rows=[
        ["Cohort Programme", "Students", "One-time / instalment", "NGN (Paystack)"],
        ["Mentorship",       "Students", "Subscription",          "NGN (Paystack)"],
        ["Consulting",       "Clients",  "Engagement / retainer", "NGN (Paystack)"],
        ["Service Request",  "Clients",  "Project-based",         "NGN (Paystack)"],
    ],
)

add_page_break(doc)

# ------------- 7. TARGET AUDIENCE -------------
add_heading(doc, "7. Target Audience & User Roles", level=1)

add_heading(doc, "7.1 Target Audience", level=2)
add_bullets(doc, [
    "Aspiring African tech talent — students and career-switchers seeking structured, "
    "affordable, mentor-led training in software, data, and product disciplines.",
    "Small and mid-sized businesses — needing reliable software, data, research, or UX "
    "work without the overhead of an in-house team.",
    "Diaspora professionals and consultants — looking for a trusted partner to staff "
    "African client engagements.",
    "Educational institutions — seeking partners for cohort delivery and mentorship.",
])

add_heading(doc, "7.2 User Roles", level=2)
add_table(
    doc,
    headers=["Role",     "Primary Surface",            "Core Capabilities"],
    rows=[
        ["Student", "/dashboard",
            "Browse + enroll in courses, attend live classes, submit assignments, track "
            "progress, manage payments."],
        ["Client",  "/client",
            "Submit service inquiries, track projects, view deliverables, manage payments."],
        ["Admin",   "/admin",
            "Manage users, courses, modules, cohorts, blogs, live sessions, service "
            "requests, assignments, payments, analytics."],
    ],
)

add_page_break(doc)

# ------------- 8. ARCHITECTURE -------------
add_heading(doc, "8. System Architecture & Technology Stack", level=1)

add_heading(doc, "8.1 Architecture Overview", level=2)
add_para(doc,
    "E-Bringgs is a modern three-tier web application: a React 18 single-page application "
    "in the browser, an Express + TypeScript REST API backed by MongoDB, and a parallel "
    "WebSocket signaling server for live video. The whole system is designed to run in "
    "three environments — development, staging, and production — each with its own "
    "MongoDB Atlas cluster.")

add_heading(doc, "8.2 Technology Stack", level=2)
add_table(
    doc,
    headers=["Layer", "Technology"],
    rows=[
        ["Frontend",   "React 18, TypeScript, Vite, Tailwind CSS v4 (CSS-first), Zustand, React Router v6, Axios"],
        ["Backend",    "Node.js, Express, TypeScript, Mongoose"],
        ["Database",   "MongoDB Atlas (per-environment clusters)"],
        ["Auth",       "JWT (access + refresh tokens), bcryptjs"],
        ["Real-time",  "WebSocket (ws library) for WebRTC signaling"],
        ["Payments",   "Paystack (primary, NGN), Stripe (legacy)"],
        ["Email",      "Nodemailer (SMTP) for verification, password reset, notifications"],
        ["AI",         "Google Gemini (Flash) for AI tutoring, with daily rate limit"],
        ["Messaging",  "Meta WhatsApp Cloud API (optional notifications)"],
        ["Security",   "Helmet, CORS, express-rate-limit, input validation, secure env handling"],
        ["Docs",       "Swagger / OpenAPI 3.0 served at /api/docs"],
    ],
    col_widths=[Inches(1.4), Inches(4.6)],
)

add_heading(doc, "8.3 Data Model (Core Entities)", level=2)
add_bullets(doc, [
    "User — name, email, password hash, role, verification state, refresh tokens.",
    "Course — title, description, modules, lessons, cohort, instructor, pricing.",
    "Blog — title, body (markdown), author, category, tags, SEO metadata.",
    "Transaction — Paystack reference, amount, currency, status, user, item.",
    "Assignment — submission file, student, course, status, reviewer feedback.",
    "ServiceInquiry — service type, client details, message, status.",
    "Project — client, scope, milestones, deliverables, status, payments.",
])

add_heading(doc, "8.4 Environments", level=2)
add_table(
    doc,
    headers=["Environment", "Purpose",                                "Database"],
    rows=[
        ["Development", "Local engineering, feature branches",   "ebringgs-development cluster"],
        ["Staging",     "QA, UAT, pre-release validation",       "ebringgs-staging cluster"],
        ["Production",  "Live customer traffic",                 "blingg-prod-cluster"],
    ],
)

add_page_break(doc)

# ------------- 9. KEY FEATURES -------------
add_heading(doc, "9. Key Features by Module", level=1)

modules = [
    ("Authentication & Identity", [
        "Email + password registration with email verification.",
        "Login / logout with role-aware redirects (admin → /admin, client → /client, student → /dashboard).",
        "Forgot password and reset password flows.",
        "JWT access tokens + refresh tokens with auto-rotation on 401.",
        "Role-based route guards on backend (protect + authorize middleware).",
    ]),
    ("Pricing & Payments", [
        "Public pricing page in NGN with responsive, animated cards.",
        "Paystack checkout (initialize + verify) integrated end-to-end.",
        "One-time and subscription billing models.",
        "Payment success and failure pages with verification.",
        "Webhook handler for asynchronous reconciliation.",
        "All transactions persisted in MongoDB.",
    ]),
    ("Services & Projects", [
        "Public services dashboard with category filtering.",
        "Service detail pages with inquiry modals.",
        "Client-side project tracking with list and detail views.",
        "Admin project management with status transitions and deliverables.",
    ]),
    ("Courses & Learning", [
        "Course catalog with search, filter, and pagination.",
        "Course detail page with modules accordion and enrollment CTA.",
        "Admin course CRUD plus per-course module editor.",
        "Cohort management for group programmes.",
        "Assignment upload and instructor review workflow.",
    ]),
    ("Live Video Classroom", [
        "Per-class WebRTC video room with join/leave.",
        "Camera and microphone toggles, screen sharing.",
        "In-class real-time chat over WebSocket.",
        "Instructor controls — mute-all, kick participants.",
        "Attendance tracking persisted server-side.",
    ]),
    ("Blog & Content", [
        "Public blog listing with search, categories, and pagination.",
        "Markdown-rendered blog posts.",
        "Admin create / edit / delete posts.",
        "SEO metadata (title, description, Open Graph) per page.",
    ]),
    ("Admin Operations", [
        "User management with role assignment.",
        "Course, module, blog, cohort, live-session management.",
        "Service request triage and assignment.",
        "Payments overview and analytics.",
        "Settings and configuration.",
    ]),
    ("Cross-Cutting UX", [
        "Light / Dark / System theme toggle (Zustand-persisted).",
        "Consistent loading, empty, and error states across pages.",
        "Floating WhatsApp contact button.",
        "Responsive layout from mobile to desktop.",
        "Accessibility basics (semantic HTML, focus styles, ARIA where needed).",
    ]),
]

for title, items in modules:
    add_heading(doc, title, level=2)
    add_bullets(doc, items)

add_page_break(doc)

# ------------- 10. MARKET LANDSCAPE -------------
add_heading(doc, "10. Market Landscape & Competitive Positioning", level=1)
add_para(doc,
    "Today, an aspiring African developer or a Nigerian SME has to navigate a fragmented "
    "landscape:")
add_bullets(doc, [
    "Foreign bootcamps (e.g. global online academies) — strong curriculum but priced in "
    "USD and disconnected from local employers.",
    "Local training schools — often physical-only, limited cohort capacity, weak "
    "post-training pipeline.",
    "Freelance marketplaces — global supply, but no training, no mentorship, and a "
    "race-to-the-bottom on pricing.",
    "Consulting firms — high quality but inaccessible to small businesses on price.",
    "DIY learning (YouTube, blogs) — free but unstructured, no accountability, no "
    "credential, no community.",
])
add_para(doc,
    "E-Bringgs sits deliberately at the intersection — close enough to each of these "
    "categories to compete, but distinct enough that no single one of them can replicate "
    "the full bundle.")

add_heading(doc, "10.1 Competitive Differentiators", level=2)
add_table(
    doc,
    headers=["Differentiator", "What It Means"],
    rows=[
        ["Local-first pricing",
            "Naira-denominated, Paystack-native — no FX friction for Nigerian users."],
        ["Three-in-one platform",
            "Training + services + consulting in one login, one dashboard."],
        ["Native live classroom",
            "WebRTC-based — no Zoom dependency, attendance baked in."],
        ["Talent-to-opportunity bridge",
            "Top students can be routed into client projects on the same platform."],
        ["Production-grade engineering",
            "TypeScript end-to-end, JWT refresh, rate limiting, Helmet, Swagger docs."],
        ["AI-assisted learning",
            "Built-in AI tutor (Gemini Flash) with daily quotas to support students."],
    ],
    col_widths=[Inches(1.7), Inches(4.3)],
)

add_page_break(doc)

# ------------- 11. CHANGING THE TREND -------------
add_heading(doc, "11. How E-Bringgs Changes the Trend", level=1)
add_para(doc,
    "The status quo treats education, services, and freelance work as three separate "
    "industries with three separate platforms. E-Bringgs argues that, for the African "
    "digital economy in particular, this separation is precisely the problem.")

add_heading(doc, "11.1 From Certificates to Careers", level=2)
add_para(doc,
    "Traditional bootcamps end on graduation day. E-Bringgs treats graduation as the "
    "midpoint — the same login that gave a student access to courses gives them visibility "
    "into client projects they can contribute to. The platform is designed so that "
    "training feeds the services arm, and the services arm validates the training.")

add_heading(doc, "11.2 From Currency Friction to Local-First", level=2)
add_para(doc,
    "By pricing in Naira and integrating Paystack natively, E-Bringgs removes the FX tax "
    "that priced thousands of capable learners out of foreign bootcamps. Pricing in NGN "
    "is not just a payments decision — it is a market-access decision.")

add_heading(doc, "11.3 From Tool Sprawl to a Single Source of Truth", level=2)
add_para(doc,
    "Schools and agencies currently run their operations across LMSs, Zoom, WhatsApp, "
    "Notion, Google Drive, and a payments tool. E-Bringgs collapses these into one "
    "system: a course is linked to a cohort, a cohort to a live session, a live session "
    "to attendance, attendance to assignments, assignments to a transcript, and the "
    "transcript to opportunities — without exporting any data.")

add_heading(doc, "11.4 From Opaque Service Delivery to Transparent Projects", level=2)
add_para(doc,
    "Clients today rarely know what state their project is in until a deadline slips. "
    "E-Bringgs gives every client a first-class dashboard — projects, deliverables, "
    "payments, and status — modelled the same way modern product teams track their own "
    "internal work. Transparency becomes a default, not a service upgrade.")

add_heading(doc, "11.5 From Passive Learning to AI-Assisted Coaching", level=2)
add_para(doc,
    "An AI tutor (Gemini Flash) is integrated into the learning experience. Students who "
    "are stuck at midnight, between live sessions, can get unblocked without waiting for "
    "the next class — closing the feedback loop that traditional cohorts cannot.")

add_heading(doc, "11.6 The Strategic Shift", level=2)
add_para(doc,
    "Taken together, E-Bringgs is repositioning the question. Instead of asking 'where "
    "do I learn to code?' or 'where do I find a developer?', users ask: 'where can I "
    "develop, deliver, and grow my career or my business in one place?' That is the "
    "trend the platform is built to set.")

add_page_break(doc)

# ------------- 12. BUSINESS MODEL -------------
add_heading(doc, "12. Business & Revenue Model", level=1)
add_table(
    doc,
    headers=["Stream", "Description", "Pricing Logic"],
    rows=[
        ["Cohort tuition",        "Structured group learning programmes",          "Fixed NGN fee per cohort, instalment supported"],
        ["Mentorship subscription","1:1 mentorship plans",                          "Recurring NGN subscription via Paystack"],
        ["Service projects",       "Software, data, research, UX delivery",        "Project-scoped quotes, milestone billing"],
        ["Consulting retainers",   "Strategic engagements",                        "Monthly retainer in NGN"],
        ["Premium content",        "Blog content, downloadable resources (future)", "Mixed free + gated"],
    ],
    col_widths=[Inches(1.5), Inches(2.5), Inches(2.0)],
)

add_page_break(doc)

# ------------- 13. IMPLEMENTATION STATUS -------------
add_heading(doc, "13. Implementation Status", level=1)
add_para(doc,
    "The platform has progressed past prototype into a near-production state. All "
    "frontend pages and backend endpoints listed above are implemented; the system is in "
    "the final pre-deployment phase.")

add_table(
    doc,
    headers=["Area", "Status"],
    rows=[
        ["Public marketing pages (Landing, Pricing, About, Contact, Terms, Privacy)", "Complete"],
        ["Authentication (register, login, refresh, verify, reset)",                  "Complete"],
        ["Student dashboard (overview, courses, resources, schedule)",                "Complete"],
        ["Client dashboard (overview, projects, payments)",                           "Complete"],
        ["Admin dashboard (users, courses, modules, blogs, payments, sessions, etc.)","Complete"],
        ["Course system (catalog, detail, modules, enrolment, assignments)",          "Complete"],
        ["Blog system (listing, post, admin CRUD)",                                   "Complete"],
        ["Services (catalog, detail, inquiry)",                                       "Complete"],
        ["Paystack payments (initialize, verify, webhook, transactions)",             "Complete"],
        ["Live video classroom (WebRTC, chat, controls, attendance)",                 "Complete"],
        ["Light / Dark / System theming",                                             "Complete"],
        ["Swagger / OpenAPI documentation",                                           "Complete"],
        [".env wiring for production secrets",                                        "Pending"],
        ["Deployment to production hosting",                                          "Pending"],
    ],
    col_widths=[Inches(4.0), Inches(1.5)],
)

add_page_break(doc)

# ------------- 14. ROADMAP -------------
add_heading(doc, "14. Future Roadmap", level=1)
add_heading(doc, "Phase 2 (Post-Launch)", level=2)
add_bullets(doc, [
    "Native mobile apps (React Native or Capacitor wrapper).",
    "Student success stories / portfolio showcase pages.",
    "Public live class schedule with countdown to next cohort.",
    "Referral / affiliate system with NGN-denominated rewards.",
    "Newsletter / email capture with exit-intent popup.",
    "Advanced analytics dashboard for admins (cohort completion, payment funnels).",
])
add_heading(doc, "Phase 3", level=2)
add_bullets(doc, [
    "Marketplace for vetted external instructors and mentors.",
    "Certificate generation with verifiable on-chain credentials (optional).",
    "Predictive learner analytics — surface at-risk students before they drop off.",
    "Expansion of payment rails into other African currencies (Flutterwave / multi-rail).",
    "Public API for institutional partners.",
])

add_page_break(doc)

# ------------- 15. RISKS -------------
add_heading(doc, "15. Risks & Mitigation", level=1)
add_table(
    doc,
    headers=["Risk", "Mitigation"],
    rows=[
        ["Payment provider downtime (Paystack)",
            "Webhook + verification fallback; Stripe code path retained as legacy."],
        ["WebRTC quality on poor networks",
            "Adaptive bitrate; chat-only fallback; recorded session playback (planned)."],
        ["Course content drift",
            "Admin module editor with versioning; cohort-scoped content references."],
        ["Security (auth, sessions, payments)",
            "JWT + refresh rotation, Helmet, rate limiting, input validation, HTTPS-only."],
        ["Operational scale (admin workload)",
            "Role assignment, scoped permissions, automation of common admin actions."],
        ["Data residency / compliance",
            "Per-environment MongoDB Atlas clusters; secrets via env; documented data model."],
    ],
    col_widths=[Inches(2.0), Inches(4.0)],
)

add_page_break(doc)

# ------------- 16. CONCLUSION -------------
add_heading(doc, "16. Conclusion", level=1)
add_para(doc,
    "E-Bringgs Technologies is not another bootcamp, not another agency, and not another "
    "freelance marketplace. It is a deliberate combination of all three — engineered as a "
    "single, role-aware product on a modern, secure, production-grade stack.")
add_para(doc,
    "By pricing in Naira, training through live mentor-led cohorts, and routing "
    "graduating talent directly into client projects, the platform addresses the three "
    "structural gaps that hold the African digital economy back: training, trust, and "
    "continuity.")
add_para(doc,
    "If the previous decade was about putting individual tools online — an LMS here, a "
    "Zoom link there, a Stripe page somewhere else — the next one is about composing "
    "them into a single experience that respects the user's time and the local "
    "economy. E-Bringgs is built for that shift.")

add_para(doc, "")
add_para(doc, "— End of Document —", italic=True, color=GREY,
         align=WD_ALIGN_PARAGRAPH.CENTER)

doc.save(OUT)
print(f"Wrote: {OUT}")
